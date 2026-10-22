"""Fluxo completo contra a operadora simulada local com modelo roteirizado."""

from __future__ import annotations

import asyncio
import csv
import json
from pathlib import Path
from typing import Any

import pytest
from faturista.__main__ import EXEMPLO, legivel
from faturista.ferramentas import FERRAMENTAS_MCP, Execucao, criar_ferramentas
from faturista.fluxo import executar, servidor_local
from faturista.planilha import ler_lote
from faturista.roteiro import ModeloRoteirizado
from langchain_core.messages import ToolMessage
from mcp.types import Tool
from openpyxl import Workbook


async def _verificar_multiglosa() -> None:
    async def chamar(_nome: str, _argumentos: dict[str, Any]) -> dict[str, Any]:
        return {}

    decisao = {
        "guia_ref": "G-1",
        "decision": "glosado",
        "motivo": "1017 DATA; 1818 COBRANÇA; 1017 REPETIDO",
        "valor_glosa_cents": 4500,
    }
    execucao = Execucao(Path("atendimentos.csv"), chamar)
    execucao.historico.extend(
        [
            ("tiss_parse_demonstrativo", {"decisoes": [decisao]}),
            ("tiss_sandbox_enviar_lote", {"aceito": True}),
        ]
    )
    ferramentas = criar_ferramentas(
        execucao,
        [
            Tool(name=nome, inputSchema={"type": "object", "properties": {}})
            for nome in FERRAMENTAS_MCP
        ],
    )
    recurso = next(tool for tool in ferramentas if tool.name == "rascunhar_recurso")
    argumentos = {
        "guia": "G-1",
        "classe": "elegibilidade",
        "justificativa": "Solicitamos revisão.",
    }
    await recurso.ainvoke({**argumentos, "codigo_glosa": "1017"})
    parcial = execucao.resultado("Parcial")
    assert parcial["status"] == "incompleto"
    assert [glosa["codigo"] for glosa in parcial["glosas"]] == ["1017", "1818"]
    assert [glosa["valor_glosa_cents"] for glosa in parcial["glosas"]] == [4500, None]
    assert "valor somado à glosa anterior" in legivel(parcial)
    await recurso.ainvoke({**argumentos, "codigo_glosa": "1818"})
    final = execucao.resultado("Completo")
    assert final["status"] == "concluido" and len(final["glosas"]) == 2

    roteiro = ModeloRoteirizado()
    mensagens = [
        ToolMessage(content='{"lote": {}}', name="ler_atendimentos", tool_call_id="lote"),
        ToolMessage(
            content=json.dumps({"decisoes": [decisao]}),
            name="tiss_parse_demonstrativo",
            tool_call_id="demonstrativo",
        ),
    ]
    explicacoes = roteiro.invoke(mensagens)
    assert {chamada["args"]["codigo"] for chamada in explicacoes.tool_calls} == {"1017", "1818"}
    mensagens.extend(
        ToolMessage(
            content=json.dumps({"codigo": codigo, "grupo": "10"}),
            name="glosa_explain",
            tool_call_id=f"glosa-{codigo}",
        )
        for codigo in ("1017", "1818")
    )
    recursos = roteiro.invoke(mensagens)
    assert {
        (chamada["args"]["guia"], chamada["args"]["codigo_glosa"])
        for chamada in recursos.tool_calls
    } == {("G-1", "1017"), ("G-1", "1818")}


def _entrada(formato: str, pasta: Path) -> Path:
    if formato == "csv":
        return EXEMPLO
    with EXEMPLO.open(encoding="utf-8", newline="") as arquivo:
        linhas = list(csv.reader(arquivo))
    destino = pasta / f"atendimentos.{formato}"
    if formato == "json":
        destino.write_text(json.dumps(ler_lote(EXEMPLO)), encoding="utf-8")
    else:
        planilha = Workbook()
        folha = planilha.active
        cabecalho = linhas[0]
        indices = {
            coluna: cabecalho.index(coluna)
            for coluna in ("numero_lote", "cnpj_prestador", "numero_guia")
        }
        folha.append(cabecalho)
        for linha in linhas[1:]:
            valores: list[Any] = list(linha)
            valores[indices["numero_lote"]] = 42
            valores[indices["cnpj_prestador"]] = 123
            valores[indices["numero_guia"]] = int(valores[indices["numero_guia"]])
            folha.append(valores)
            for coluna, formato_celula in (
                ("numero_lote", "0000"),
                ("cnpj_prestador", "0" * 14),
                ("numero_guia", "0000"),
            ):
                folha.cell(
                    row=folha.max_row, column=indices[coluna] + 1
                ).number_format = formato_celula
        planilha.save(destino)
    return destino


@pytest.mark.parametrize(
    ("formato", "cenario", "glosas"),
    [
        ("csv", "carteirinha_vencida", {("0001", "1017"), ("0002", "1017"), ("0003", "1017")}),
        ("xlsx", "sem_autorizacao", {("0001", "1818")}),
        ("json", "aprovado", set()),
    ],
)
def test_fluxo_fecha_o_ciclo(
    formato: str, cenario: str, glosas: set[tuple[str, str]], tmp_path: Path
) -> None:
    arquivo = _entrada(formato, tmp_path)
    if formato == "xlsx":
        lote_entrada = ler_lote(arquivo)
        assert lote_entrada["numero_lote"] == "0042"
        assert lote_entrada["prestador"]["cnpj"] == "00000000000123"
        assert lote_entrada["guias"][0]["cabecalho_guia"]["numero_guia_prestador"] == "0001"
        assert lote_entrada["guias"][0]["dados_beneficiario"]["numero_carteira"] == "12345"
        assert (
            lote_entrada["guias"][0]["dados_solicitante"]["nome_contratado_solicitante"]
            == "CLINICA EXEMPLO"
        )
    with servidor_local() as url:
        resultado = asyncio.run(
            executar(arquivo, ModeloRoteirizado(cenario=cenario), url=url, cenario=cenario)
        )
    if cenario == "aprovado":
        asyncio.run(_verificar_multiglosa())
    lote = resultado["lote"]
    assert (resultado["status"], lote["aceito"], lote["valido"]) == ("concluido", True, True)
    assert lote["numero_protocolo"] and lote["rodadas_de_correcao"] == 1
    assert {(g["guia"], g["codigo"]) for g in resultado["glosas"]} == glosas
    assert {(r["guia"], r["codigo_glosa"]) for r in resultado["recursos"]} == glosas
    for recurso in resultado["recursos"]:
        assert recurso["needs"] == ["evidencia_clinica"]
        assert recurso["cta"]["url"] == "https://tiss-api.voa.health/recurso"
    assert "aceito" in resultado["resumo"] and lote["numero_lote"] in legivel(resultado)
