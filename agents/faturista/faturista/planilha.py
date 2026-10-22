"""Leitura de atendimentos em CSV, XLSX ou JSON para um lote TISS em JSON."""

from __future__ import annotations

import csv
import json
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _centavos(valor: Any) -> int:
    texto = str(valor).strip()
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    return int((Decimal(texto) * 100).quantize(Decimal("1")))


def _linhas(caminho: Path) -> list[dict[str, Any]]:
    sufixo = caminho.suffix.lower()
    if sufixo == ".csv":
        with caminho.open(encoding="utf-8", newline="") as arquivo:
            return list(csv.DictReader(arquivo))
    if sufixo == ".xlsx":
        planilha = load_workbook(caminho, read_only=True, data_only=True).active
        if planilha is None:
            return []
        linhas = [list(linha) for linha in planilha.iter_rows(values_only=True)]
        cabecalho = [str(coluna) for coluna in linhas[0]]
        return [dict(zip(cabecalho, linha, strict=False)) for linha in linhas[1:] if any(linha)]
    if sufixo == ".json":
        dados = json.loads(caminho.read_text(encoding="utf-8"))
        return dados if isinstance(dados, list) else [dados]
    raise ValueError(f"Formato não suportado: {caminho.suffix}")


def _profissional(linha: dict[str, Any]) -> dict[str, str]:
    return {
        "cbos": str(linha["cbos"]),
        "conselho_profissional": "01",
        "nome_profissional": str(linha["profissional"]),
        "numero_conselho_profissional": str(linha["crm"]),
        "uf": str(linha["uf_conselho"]),
    }


def _guia(numero: str, linhas: list[dict[str, Any]]) -> dict[str, Any]:
    primeira = linhas[0]
    contratado = {"codigo_prestador_na_operadora": str(primeira["codigo_prestador"])}
    itens = [
        {
            "sequencial_item": str(indice),
            "data_execucao": str(linha["data"]),
            "procedimento": {
                "codigo_tabela": str(linha["codigo_tabela"]),
                "codigo_procedimento": str(linha["codigo_procedimento"]),
                "descricao_procedimento": str(linha["descricao"]),
            },
            "quantidade_executada": str(linha["quantidade"]),
            "reducao_acrescimo": "1.00",
            "valor_unitario": _centavos(linha["valor_unitario"]),
            "valor_total": _centavos(linha["valor_total"]),
        }
        for indice, linha in enumerate(linhas, start=1)
    ]
    total = sum(_centavos(linha["valor_total"]) for linha in linhas)
    return {
        "tipo": "sp_sadt",
        "cabecalho_guia": {
            "numero_guia_prestador": numero,
            "registro_ans": str(primeira["registro_ans"]),
        },
        "dados_beneficiario": {
            "atendimento_rn": "N",
            "numero_carteira": str(primeira["numero_carteira"]),
        },
        "dados_solicitante": {
            "contratado_solicitante": contratado,
            "nome_contratado_solicitante": "CLINICA EXEMPLO",
            "profissional_solicitante": _profissional(primeira),
        },
        "dados_solicitacao": {"carater_atendimento": "1"},
        "dados_executante": {"cnes": str(primeira["cnes"]), "contratado_executante": contratado},
        "dados_atendimento": {
            "indicacao_acidente": "0",
            "regime_atendimento": "01",
            "tipo_atendimento": "01",
        },
        "procedimentos_executados": itens,
        "valor_total": {"valor_procedimentos": total, "valor_total_geral": total},
    }


def ler_lote(caminho: str | Path) -> dict[str, Any]:
    """Converta atendimentos, uma linha por procedimento, em lote TISS."""
    caminho = Path(caminho)
    if caminho.suffix.lower() == ".json":
        dados = json.loads(caminho.read_text(encoding="utf-8"))
        if isinstance(dados, dict) and "versao" in dados:
            return dados
    linhas = _linhas(caminho)
    if not linhas:
        raise ValueError("Planilha sem atendimentos")
    guias: dict[str, list[dict[str, Any]]] = {}
    for linha in linhas:
        guias.setdefault(str(linha["numero_guia"]), []).append(linha)
    primeira = linhas[0]
    return {
        "versao": "4.03.00",
        "transacao": {
            "sequencial_transacao": str(primeira["numero_lote"]).lstrip("0") or "1",
            "data_registro_transacao": max(str(linha["data"]) for linha in linhas),
            "hora_registro_transacao": "12:00:00",
        },
        "prestador": {"cnpj": str(primeira["cnpj_prestador"])},
        "operadora": {"registro_ans": str(primeira["registro_ans"])},
        "numero_lote": str(primeira["numero_lote"]),
        "guias": [_guia(numero, itens) for numero, itens in guias.items()],
    }
